"""
Importa uma turma de projeto (ex: "3167035 - LE - APRENDER E ENSINAR / PAP
/ PAP COLABORATIVO 3 4 E 5 ANO") a partir do relatório "Alunos da Turma"
do SGP.

O arquivo que o SGP exporta tem extensão .xls, mas por dentro é uma
tabela HTML. Também aceita o mesmo relatório salvo como .xlsx. Formato:

    Unidade   | 000000 - EMEF - ESCOLA MODELO (FICTÍCIA)
    Turma     | 3167035 - LE | Situação | Organizada | Ano Letivo | 2026 | Turno/Duração | Manhã 01:30h
    Descrição | APRENDER E ENSINAR / PAP / PAP COLABORATIVO 3 4 E 5 ANO
    (linha em branco)
    Código | RA Prodesp | Código INEP | Nº | Nome do Aluno | ... | Situação Aluno | ... | Série | ...
    ...uma linha por aluno...

Diferente da lista PILOTO, isso não muda a turma regular do aluno: os
alunos de projeto também estão numa turma regular. A turma de projeto
fica nas tabelas turmas_projeto e turma_projeto_alunos, e cada linha é
ligada ao cadastro pelo código EOL (coluna "Código"). Se o aluno não for
encontrado, a linha fica guardada só com o nome.

Nome da turma = "<Turma> - <Descrição>"; turno = "Turno/Duração" sem a
duração ("Manhã 01:30h" -> "Manhã").

Importar de novo a mesma turma (mesmo código) atualiza o cabeçalho e
substitui a lista de alunos.
"""
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from models import agora, get_db


class _Tabela(HTMLParser):
    """Linhas (listas de textos das células) de todas as tabelas do HTML."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.linhas, self._linha, self._celula = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._linha = []
        elif tag in ("td", "th") and self._linha is not None:
            self._celula = ""

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._linha is not None and self._celula is not None:
            self._linha.append(" ".join(self._celula.split()))
            self._celula = None
        elif tag == "tr" and self._linha is not None:
            self.linhas.append(self._linha)
            self._linha = None

    def handle_data(self, data):
        if self._celula is not None:
            self._celula += data


def _linhas_do_arquivo(caminho):
    caminho = Path(caminho)
    bruto = caminho.read_bytes()
    if bruto[:4] == b"PK\x03\x04":  # .xlsx
        from openpyxl import load_workbook
        wb = load_workbook(caminho, read_only=True, data_only=True)
        ws = wb.active
        linhas = [[("" if v is None else " ".join(str(v).split())) for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return linhas
    if bruto[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":  # .xls binário de verdade
        raise ValueError('Arquivo .xls no formato antigo do Excel: abra e use "Salvar como" .xlsx antes de enviar.')
    for codificacao in ("utf-8", "cp1252", "latin-1"):
        try:
            texto = bruto.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    if "<t" not in texto.lower():
        raise ValueError("O arquivo não parece o relatório \"Alunos da Turma\" do SGP.")
    parser = _Tabela()
    parser.feed(texto)
    return parser.linhas


def _valor_ao_lado(linha, rotulo):
    """Texto da célula logo depois da célula com esse rótulo (sem acento/caixa)."""
    alvo = _chave(rotulo)
    for i, cel in enumerate(linha[:-1]):
        if _chave(cel) == alvo:
            return linha[i + 1].strip()
    return ""


def _chave(texto):
    import unicodedata
    t = unicodedata.normalize("NFD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9/]", "", t.lower())


def turno_sem_duracao(texto):
    """'Manhã 01:30h' -> 'Manhã'."""
    return re.sub(r"\s*\d{1,2}[:h]\d{2}\s*h?\s*$", "", (texto or "").strip(), flags=re.I).strip()


def ler_arquivo(caminho):
    """Devolve {"turma": {...cabeçalho...}, "alunos": [{...}]}."""
    linhas = [l for l in _linhas_do_arquivo(caminho)]
    cab = {"unidade": "", "turma": "", "descricao": "", "turno": "", "ano": None}
    idx_cabecalho_alunos = None
    for i, linha in enumerate(linhas):
        chaves = [_chave(c) for c in linha]
        if "unidade" in chaves and not cab["unidade"]:
            cab["unidade"] = _valor_ao_lado(linha, "Unidade")
        if "turma" in chaves and not cab["turma"]:
            cab["turma"] = _valor_ao_lado(linha, "Turma")
            cab["turno"] = turno_sem_duracao(_valor_ao_lado(linha, "Turno/Duração"))
            ano = _valor_ao_lado(linha, "Ano Letivo")
            cab["ano"] = int(ano) if ano.isdigit() else None
        if "descricao" in chaves and not cab["descricao"]:
            cab["descricao"] = _valor_ao_lado(linha, "Descrição")
        if "codigo" in chaves and "nomedoaluno" in chaves:
            idx_cabecalho_alunos = i
            break

    if not cab["turma"] or idx_cabecalho_alunos is None:
        raise ValueError(
            "Não encontrei a turma e a lista de alunos. Envie o relatório \"Alunos da Turma\" do SGP "
            "(com as linhas Turma, Descrição e Turno/Duração no topo)."
        )

    colunas = {_chave(c): j for j, c in enumerate(linhas[idx_cabecalho_alunos])}

    def col(linha, *nomes):
        for nome in nomes:
            j = colunas.get(_chave(nome))
            if j is not None and j < len(linha):
                return linha[j].strip()
        return ""

    alunos = []
    for linha in linhas[idx_cabecalho_alunos + 1:]:
        codigo, nome = col(linha, "Código"), col(linha, "Nome do Aluno")
        if not codigo or not nome:
            continue
        alunos.append({
            "codigo": re.sub(r"\D", "", codigo) or codigo,
            "numero": col(linha, "Nº", "N"),
            "nome": nome,
            "situacao": col(linha, "Situação Aluno") or "—",
            "data_inclusao": col(linha, "Data Inclusão Aluno Turma"),
        })

    turma_txt = cab["turma"]
    codigo_turma = turma_txt.split("-")[0].strip()
    nome = f"{turma_txt} - {cab['descricao']}" if cab["descricao"] else turma_txt
    return {"turma": {**cab, "codigo": codigo_turma, "nome": nome}, "alunos": alunos}


def importar(caminho):
    dados = ler_arquivo(caminho)
    t = dados["turma"]
    with get_db() as db:
        existente = db.execute("SELECT id FROM turmas_projeto WHERE codigo = ?", (t["codigo"],)).fetchone()
        if existente:
            turma_id = existente["id"]
            db.execute(
                """UPDATE turmas_projeto SET nome=?, turma=?, descricao=?, turno=?, ano=?, unidade=?, updated_at=?
                   WHERE id=?""",
                (t["nome"], t["turma"], t["descricao"], t["turno"], t["ano"], t["unidade"], agora(), turma_id),
            )
            db.execute("DELETE FROM turma_projeto_alunos WHERE turma_projeto_id = ?", (turma_id,))
        else:
            turma_id = db.execute(
                """INSERT INTO turmas_projeto (codigo, nome, turma, descricao, turno, ano, unidade, updated_at)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (t["codigo"], t["nome"], t["turma"], t["descricao"], t["turno"], t["ano"], t["unidade"], agora()),
            ).lastrowid

        ligados, nao_encontrados, ativos = 0, [], 0
        vistos = set()
        for a in dados["alunos"]:
            if a["codigo"] in vistos:
                continue
            vistos.add(a["codigo"])
            aluno = db.execute("SELECT id FROM students WHERE codigo_interno = ?", (a["codigo"],)).fetchone()
            db.execute(
                """INSERT INTO turma_projeto_alunos
                   (turma_projeto_id, student_id, codigo_eol, numero, nome, situacao, data_inclusao)
                   VALUES (?,?,?,?,?,?,?)""",
                (turma_id, aluno["id"] if aluno else None, a["codigo"], a["numero"], a["nome"],
                 a["situacao"], a["data_inclusao"]),
            )
            if aluno:
                ligados += 1
            else:
                nao_encontrados.append(a["nome"])
            if _chave(a["situacao"]) == "ativo":
                ativos += 1

    return {
        "turma_id": turma_id, "nome": t["nome"], "turno": t["turno"], "novo": not existente,
        "total": len(vistos), "ativos": ativos, "ligados": ligados, "nao_encontrados": nao_encontrados,
    }


if __name__ == "__main__":
    r = importar(sys.argv[1])
    print(f"{r['nome']} ({r['turno']}): {r['total']} aluno(s), {r['ativos']} ativo(s), "
          f"{r['ligados']} ligado(s) ao cadastro")
